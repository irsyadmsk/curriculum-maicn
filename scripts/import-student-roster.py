"""Read the supplied placement workbook; export names and existing classes only.

Run with the bundled Python runtime. The workbook is never edited. Marks,
placement rankings and private notes are not copied into the portal.
"""
from pathlib import Path
import argparse
import hashlib
import json
import re
import unicodedata

import openpyxl


def student_id(name):
    key = unicodedata.normalize('NFC', name).strip().casefold()
    return 'murid-' + hashlib.sha256(('maicn:2026-2027:' + key).encode()).hexdigest()[:16]


def extract(path):
    original_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    summary = workbook['Ringkasan & Catatan']
    counts = {row[0]: row[1] for row in summary.iter_rows(values_only=True)
              if row[0] in ('X-1', 'X-2', 'X-3') and isinstance(row[1], int)}
    classes, references, seen = {}, {}, set()
    for classroom in ('X-1', 'X-2', 'X-3'):
        sheet = workbook[classroom]
        if sheet['A4'].value != 'No' or sheet['B4'].value != 'Nama Murid':
            raise ValueError(f'Unexpected name columns in {classroom}')
        students = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), 5):
            ordinal, name = row[:2]
            if not isinstance(ordinal, int):
                continue  # header/total, never interpreted as a pupil
            if not isinstance(name, str) or not name.strip():
                raise ValueError(f'Missing name at {classroom}!B{row_number}')
            identifier = student_id(name)
            if identifier in seen:
                raise ValueError(f'Duplicate pupil at {classroom}!B{row_number}')
            if ordinal != len(students) + 1:
                raise ValueError(f'Unexpected source ordinal at {classroom}!A{row_number}')
            seen.add(identifier)
            students.append({'id': identifier, 'nama': name, 'kelas': classroom,
                             'nomor': ordinal, 'nis': None, 'nisn': None,
                             'source': f'{classroom}!B{row_number}'})
        if len(students) != counts.get(classroom):
            raise ValueError(f'Class count does not match workbook summary: {classroom}')
        classes[classroom] = students
        references[classroom] = f'B5:B{4 + len(students)}'

    pending = []
    # The workbook mentions one pupil outside the class sheets. This is data
    # about an absent class assignment, never an instruction to re-rank pupils.
    for row_number, row in enumerate(summary.iter_rows(values_only=True), 1):
        text = row[0]
        if not isinstance(text, str):
            continue
        match = re.match(r'^(.+?) tidak memiliki nilai sama sekali', text)
        if match:
            name = match.group(1)
            identifier = student_id(name)
            if identifier not in seen:
                pending.append({'id': identifier, 'nama': name, 'kelas': None,
                                'nomor': None, 'nis': None, 'nisn': None,
                                'source': f'Ringkasan & Catatan!A{row_number}'})
                seen.add(identifier)
    workbook.close()
    if hashlib.sha256(path.read_bytes()).hexdigest() != original_hash:
        raise ValueError('Source changed during extraction')
    return {'schemaVersion': 1, 'tahunPelajaran': '2026/2027',
            'source': {'file': path.name, 'sha256': original_hash, 'ranges': references},
            'classes': classes, 'unassigned': pending}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('--output', type=Path, default=Path(__file__).resolve().parents[1] / 'assets/student-roster-data.js')
    args = parser.parse_args()
    data = extract(args.source)
    args.output.write_text('// Local roster imported from the supplied workbook. No grades or private placement notes.\n'
                           'window.MAICN_STUDENT_ROSTER = ' + json.dumps(data, ensure_ascii=False, indent=2) + ';\n', encoding='utf-8')
    summary = {'assigned': sum(len(rows) for rows in data['classes'].values()), 'unassigned': len(data['unassigned']), 'classes': {key: len(rows) for key, rows in data['classes'].items()}}
    args.output.with_name('student-roster-summary.js').write_text('window.MAICN_ROSTER_SUMMARY = Object.freeze(' + json.dumps(summary) + ');\n', encoding='utf-8')
    print(json.dumps({'classes': {key: len(value) for key, value in data['classes'].items()},
                      'unassigned': len(data['unassigned']), 'sourceSHA256': data['source']['sha256']}))
